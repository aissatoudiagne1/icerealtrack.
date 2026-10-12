"""Génération du rapport Word de I.CerealTrack."""
import base64
import io
from datetime import date

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Seuils des recommandations : à valider avec ton encadrant.
SEUIL_PH_MIN = 6.0
SEUIL_PH_MAX = 7.8
SEUIL_TEMP_STRESS = 35.0      # °C
SEUIL_ECART_RELATIF = 0.25    # 25 % d'écart par rapport à la moyenne des sites

VARS_SOL = ["ph", "conductivite", "salinite", "n", "p", "k", "temper", "humidite"]
LABELS_SOL = {
    "ph": "pH", "conductivite": "Conductivité", "salinite": "Salinité",
    "n": "Azote (N)", "p": "Phosphore (P)", "k": "Potassium (K)",
    "temper": "Temp. sol", "humidite": "Humidité sol",
}
VARS_CLIMAT = ["temp_min", "temp_max", "temp_moy", "humidite_moy",
               "precipitation", "vent", "rayonnement"]
LABELS_CLIMAT = {
    "temp_min": "T° min", "temp_max": "T° max", "temp_moy": "T° moy",
    "humidite_moy": "Humidité moy", "precipitation": "Pluie (mm)",
    "vent": "Vent (m/s)", "rayonnement": "Rayonnement",
}


# ---------- Outils Word ----------
def _tableau(doc, df):
    if df is None or df.empty:
        doc.add_paragraph("Aucune donnée disponible.")
        return
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"
    for i, c in enumerate(df.columns):
        cell = t.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(str(c)).bold = True
    for _, ligne in df.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(ligne):
            if pd.isna(v):
                s = "—"
            elif isinstance(v, float):
                s = f"{v:.2f}"
            else:
                s = str(v)
            cells[i].text = s
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()


def _ajouter_figure(doc, fig, largeur=15.5):
    tampon = io.BytesIO()
    fig.savefig(tampon, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    tampon.seek(0)
    doc.add_picture(tampon, width=Cm(largeur))


# ---------- Statistiques ----------
def _resume_moyennes(df, variables, labels):
    lignes = []
    for site, g in df.groupby("nomsite"):
        ligne = {"Site": site, "N": len(g)}
        for v in variables:
            s = g[v].dropna()
            if len(s) > 1:
                ligne[labels[v]] = f"{s.mean():.2f} ({s.std():.2f})"
            elif len(s) == 1:
                ligne[labels[v]] = f"{s.mean():.2f}"
            else:
                ligne[labels[v]] = "—"
        lignes.append(ligne)
    return pd.DataFrame(lignes)


def _extremes(df, var, label):
    m = df.groupby("nomsite")[var].mean().dropna()
    if len(m) < 2:
        return None
    return (f"{label} : moyenne la plus élevée à {m.idxmax()} ({m.max():.2f}), "
            f"la plus faible à {m.idxmin()} ({m.min():.2f}).")


def _ecarts_relatifs(mesures, var):
    moy = mesures.groupby("nomsite")[var].mean().dropna()
    if len(moy) < 2:
        return {}
    ref = moy.mean()
    if ref == 0:
        return {}
    return {s: (v - ref) / abs(ref) for s, v in moy.items()}


def _varietes_resistantes(varietes, colonne):
    if colonne not in varietes.columns:
        return []
    ok = varietes[varietes[colonne].astype(str).str.strip().str.lower() == "resistante"]
    return ok["nomvariete"].dropna().tolist()


# ---------- Graphiques ----------
def _fig_boites(mesures, variables):
    sites = sorted(mesures["nomsite"].dropna().unique())
    fig, axes = plt.subplots(1, len(variables), figsize=(4.2 * len(variables), 3.8))
    if len(variables) == 1:
        axes = [axes]
    for ax, v in zip(axes, variables):
        sites_v = [s for s in sites if mesures.loc[mesures["nomsite"] == s, v].notna().any()]
        if not sites_v:
            ax.set_visible(False)
            continue
        ax.boxplot([mesures.loc[mesures["nomsite"] == s, v].dropna() for s in sites_v])
        ax.set_xticks(range(1, len(sites_v) + 1))
        ax.set_xticklabels(sites_v, rotation=20)
        ax.set_title(LABELS_SOL[v])
    fig.tight_layout()
    return fig


def _fig_nutriments(mesures):
    moy = mesures.groupby("nomsite")[["n", "p", "k"]].mean().dropna(how="all")
    if moy.empty:
        return None
    fig, ax = plt.subplots(figsize=(7, 3.8))
    moy.rename(columns={"n": "Azote (N)", "p": "Phosphore (P)", "k": "Potassium (K)"}).plot(kind="bar", ax=ax)
    ax.set_xlabel("")
    ax.set_ylabel("Valeur moyenne")
    ax.tick_params(axis="x", rotation=20)
    fig.tight_layout()
    return fig


def _fig_climat(climat):
    c = climat.copy()
    c["datemesure"] = pd.to_datetime(c["datemesure"])
    fig, axes = plt.subplots(2, 1, figsize=(8, 6), sharex=True)
    for site, g in c.groupby("nomsite"):
        g = g.sort_values("datemesure")
        axes[0].plot(g["datemesure"], g["temp_moy"], label=site)
        axes[1].plot(g["datemesure"], g["humidite_moy"], label=site)
    axes[0].set_ylabel("Température moy. (°C)")
    axes[1].set_ylabel("Humidité moy. (%)")
    axes[0].legend()
    fig.autofmt_xdate()
    fig.tight_layout()
    return fig


def _fig_pluie(climat):
    c = climat.copy()
    c["mois"] = pd.to_datetime(c["datemesure"]).dt.to_period("M").astype(str)
    p = c.groupby(["mois", "nomsite"])["precipitation"].sum().unstack()
    if p.empty or p.dropna(how="all").empty:
        return None
    fig, ax = plt.subplots(figsize=(8, 3.6))
    p.plot(kind="bar", ax=ax)
    ax.set_xlabel("")
    ax.set_ylabel("Précipitations (mm)")
    ax.tick_params(axis="x", rotation=30)
    fig.tight_layout()
    return fig


# ---------- Recommandations ----------
def recommandations(mesures, climat, synthese, varietes):
    recs = []

    for site, g in mesures.groupby("nomsite"):
        ph = g["ph"].dropna()
        if ph.empty:
            continue
        m = ph.mean()
        if m < SEUIL_PH_MIN:
            recs.append(("Attention", f"{site} : pH moyen de {m:.2f}, sol acide (seuil {SEUIL_PH_MIN}). "
                         "À confirmer par une analyse de laboratoire ; un amendement calcique peut être envisagé."))
        elif m > SEUIL_PH_MAX:
            recs.append(("Attention", f"{site} : pH moyen de {m:.2f}, sol alcalin (seuil {SEUIL_PH_MAX}). "
                         "Surveiller la disponibilité du phosphore et des oligo-éléments ; un apport de matière organique peut aider."))
        else:
            recs.append(("OK", f"{site} : pH moyen de {m:.2f}, dans la plage favorable au blé "
                         f"({SEUIL_PH_MIN} à {SEUIL_PH_MAX})."))

    sel_haute = False
    for var, label in [("salinite", "salinité"), ("conductivite", "conductivité électrique")]:
        for site, e in _ecarts_relatifs(mesures, var).items():
            if e > SEUIL_ECART_RELATIF:
                sel_haute = True
                recs.append(("Attention", f"{site} : {label} supérieure de {e * 100:.0f} % à la moyenne des sites. "
                             "Prévoir un suivi renforcé, un drainage ou un lessivage adapté, et privilégier des variétés tolérantes au sel."))
    if sel_haute:
        tol = _varietes_resistantes(varietes, "resistance_salinite")
        if tol:
            recs.append(("Info", "Variétés de la base notées résistantes à la salinité : " + ", ".join(tol) + "."))

    for var, label in [("n", "azote (N)"), ("p", "phosphore (P)"), ("k", "potassium (K)")]:
        for site, e in _ecarts_relatifs(mesures, var).items():
            if e < -SEUIL_ECART_RELATIF:
                recs.append(("Attention", f"{site} : {label} inférieur de {abs(e) * 100:.0f} % à la moyenne des sites. "
                             "Un apport de fertilisation ciblé est à étudier avec l'agronome."))

    chaleur = False
    for site, g in climat.groupby("nomsite"):
        tmax = g["temp_max"].dropna()
        if tmax.empty:
            continue
        n = int((tmax > SEUIL_TEMP_STRESS).sum())
        if n:
            chaleur = True
            recs.append(("Attention", f"{site} : {n} jour(s) sur {len(tmax)} avec une température maximale "
                         f"supérieure à {SEUIL_TEMP_STRESS:.0f} °C (maximum {tmax.max():.1f} °C). "
                         "Risque de stress thermique, surtout à la floraison : adapter la date de semis et choisir des variétés tolérantes à la chaleur."))
    if chaleur:
        tol = _varietes_resistantes(varietes, "resistance_chaleur")
        if tol:
            recs.append(("Info", "Variétés de la base notées résistantes à la chaleur : " + ", ".join(tol) + "."))

    if "rendement_moyen" in synthese.columns and "rendement_moyen_isra" in varietes.columns:
        comp = synthese.merge(varietes[["nomvariete", "rendement_moyen_isra"]], on="nomvariete", how="left")
        for _, r in comp.iterrows():
            rm, ref = r.get("rendement_moyen"), r.get("rendement_moyen_isra")
            if pd.notna(rm) and pd.notna(ref) and float(ref) > 0:
                ecart = (float(rm) / float(ref) - 1) * 100
                txt = (f"Essai {r['idexper']} ({r['nomvariete']}, {r['nomsite']}) : rendement mesuré {float(rm):.2f} "
                       f"contre {float(ref):.2f} en référence ISRA ({ecart:+.0f} %). "
                       "Comparaison valable si les unités sont identiques (T/ha).")
                if ecart < -20:
                    recs.append(("Attention", txt + " Écart important à analyser (sol, climat, conduite de l'essai)."))
                elif ecart > 10:
                    recs.append(("Positif", txt))
                else:
                    recs.append(("Info", txt))

    if not recs:
        recs.append(("Info", "Pas assez de données pour formuler des recommandations."))
    return recs


def _champ_page(paragraphe):
    run = paragraphe.add_run()
    debut = OxmlElement("w:fldChar")
    debut.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    run._r.append(debut)
    run._r.append(instr)
    run._r.append(fin)


# ---------- Rapport ----------
def generer_rapport(mesures, climat, synthese, varietes, sites_choisis, periode, utilisateur, logo_b64=None):
    mesures = mesures.copy()
    climat = climat.copy()
    for v in VARS_SOL:
        if v in mesures.columns:
            mesures[v] = pd.to_numeric(mesures[v], errors="coerce")
    for v in VARS_CLIMAT:
        if v in climat.columns:
            climat[v] = pd.to_numeric(climat[v], errors="coerce")
    synthese = synthese[synthese["nomsite"].isin(sites_choisis)].copy()

    if len(periode) == 2:
        txt_periode = f"du {periode[0]:%d/%m/%Y} au {periode[1]:%d/%m/%Y}"
    else:
        txt_periode = "toute la période disponible"

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    # --- Page de garde ---
    if logo_b64:
        try:
            brut = logo_b64.split("base64,")[-1]
            doc.add_picture(io.BytesIO(base64.b64decode(brut)), width=Cm(4))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            pass
    for _ in range(3):
        doc.add_paragraph()
    titre = doc.add_paragraph()
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t1 = titre.add_run("I.Cereal")
    t1.bold = True
    t1.font.size = Pt(40)
    t1.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)
    t2 = titre.add_run("Track")
    t2.bold = True
    t2.font.size = Pt(40)
    t2.font.color.rgb = RGBColor(0xD4, 0xAF, 0x37)
    sous = doc.add_paragraph()
    sous.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sous.add_run("Rapport d'analyse – suivi agronomique du blé").font.size = Pt(18)
    inst = doc.add_paragraph()
    inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inst.add_run("ISRA – CRA Saint-Louis").font.size = Pt(14)
    for _ in range(3):
        doc.add_paragraph()
    for ligne in (f"Sites : {', '.join(sites_choisis) or '—'}",
                  f"Période : {txt_periode}",
                  f"Généré le {date.today():%d/%m/%Y} par {utilisateur['prenom']} {utilisateur['nom']}"):
        p = doc.add_paragraph(ligne)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    section = doc.sections[0]
    section.different_first_page_header_footer = True
    pied = section.footer.paragraphs[0]
    pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pied.add_run("I.CerealTrack · ISRA CRA Saint-Louis · page ")
    _champ_page(pied)

    doc.add_heading("1. Résumé chiffré", 1)
    nb_points = mesures["idpoint"].nunique() if "idpoint" in mesures.columns else "—"
    _tableau(doc, pd.DataFrame([
        {"Indicateur": "Sites analysés", "Valeur": mesures["nomsite"].nunique()},
        {"Indicateur": "Essais", "Valeur": len(synthese)},
        {"Indicateur": "Points d'échantillonnage", "Valeur": nb_points},
        {"Indicateur": "Mesures du sol", "Valeur": len(mesures)},
        {"Indicateur": "Relevés climatiques", "Valeur": len(climat)},
    ]))

    doc.add_heading("2. Analyse du sol", 1)
    if mesures.empty:
        doc.add_paragraph("Aucune mesure du sol sur la sélection.")
    else:
        doc.add_paragraph("Moyenne (écart-type) par site :")
        _tableau(doc, _resume_moyennes(mesures, VARS_SOL, LABELS_SOL))
        _ajouter_figure(doc, _fig_boites(mesures, ["ph", "conductivite", "salinite"]))
        fig = _fig_nutriments(mesures)
        if fig is not None:
            _ajouter_figure(doc, fig, 13)
        for var in ["ph", "conductivite", "salinite", "n", "p", "k"]:
            phrase = _extremes(mesures, var, LABELS_SOL[var])
            if phrase:
                doc.add_paragraph(phrase, style="List Bullet")

    doc.add_heading("3. Analyse du climat", 1)
    if climat.empty:
        doc.add_paragraph("Aucun relevé climatique sur la sélection.")
    else:
        doc.add_paragraph("Moyenne (écart-type) par site :")
        _tableau(doc, _resume_moyennes(climat, VARS_CLIMAT, LABELS_CLIMAT))
        _ajouter_figure(doc, _fig_climat(climat))
        fig = _fig_pluie(climat)
        if fig is not None:
            _ajouter_figure(doc, fig)
        for var in ["temp_moy", "humidite_moy", "vent"]:
            phrase = _extremes(climat, var, LABELS_CLIMAT[var])
            if phrase:
                doc.add_paragraph(phrase, style="List Bullet")

    doc.add_heading("4. Essais et rendement", 1)
    colonnes = {"idexper": "Essai", "nomsite": "Site", "nomvariete": "Variété",
                "nb_points": "Points", "nb_mesures": "Mesures",
                "ph_moyen": "pH moy.", "rendement_moyen": "Rendement moy."}
    dispo = [c for c in colonnes if c in synthese.columns]
    _tableau(doc, synthese[dispo].rename(columns=colonnes))

    doc.add_heading("5. Qualité des données", 1)
    if mesures.empty:
        doc.add_paragraph("Sans objet.")
    else:
        manquants = [{"Variable": LABELS_SOL[v], "Valeurs manquantes (%)": float(mesures[v].isna().mean() * 100)}
                     for v in VARS_SOL if v in mesures.columns]
        _tableau(doc, pd.DataFrame(manquants))

    doc.add_heading("6. Observations et recommandations", 1)
    for niveau, texte in recommandations(mesures, climat, synthese, varietes):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"[{niveau}] ").bold = True
        p.add_run(texte)

    doc.add_heading("7. Limites et méthode", 1)
    doc.add_paragraph(
        "Les statistiques sont calculées sur les données de la plateforme pour les sites et la période choisis. "
        "Les recommandations sont produites automatiquement par des règles à seuils (pH, température maximale, "
        "écarts relatifs entre sites, comparaison à la référence ISRA). Elles constituent une aide à l'analyse "
        "et doivent être validées par un agronome avant toute décision.")

    tampon = io.BytesIO()
    doc.save(tampon)
    return tampon.getvalue()
